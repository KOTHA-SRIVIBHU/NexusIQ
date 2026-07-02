from pptx import Presentation
from typing import Dict


def parse_pptx(filepath: str) -> Dict:
    prs = Presentation(filepath)
    slides = []
    for i, slide in enumerate(prs.slides):
        texts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        slides.append({"slide": i + 1, "text": "\n".join(texts)})
    full_text = "\n".join(s.text for s in slides)
    return {
        "text": full_text,
        "pages": slides,
        "page_count": len(slides),
    }
